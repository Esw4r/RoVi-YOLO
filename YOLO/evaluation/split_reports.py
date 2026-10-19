"""Create four project-specific Word reports from the measured evaluation."""
from pathlib import Path
import json, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
E=ROOT/'output/evaluation'; OUT=ROOT/'output/split_reports'; OUT.mkdir(exist_ok=True)
def load(n): return json.loads((E/n).read_text())
A=load('audit.json'); B=load('baseline_metrics.json'); D=load('analysis.json'); X=load('extended_results.json'); M=load('model_metadata.json')
def pc(n): return f'{100*n:.2f}%'
def add(d,t,style=None): return d.add_paragraph(t,style)
def head(d,t): d.add_heading(t,1)
def sub(d,t): d.add_heading(t,2)
def page(d,t): d.add_page_break(); head(d,t)
def new(title,scope):
    d=Document(); sec=d.sections[0]
    sec.page_width=Inches(8.27); sec.page_height=Inches(11.69)
    sec.top_margin=sec.bottom_margin=Inches(.7); sec.left_margin=sec.right_margin=Inches(.75)
    normal=d.styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(7); normal.paragraph_format.line_spacing=1.12
    for n,size in [('Title',23),('Subtitle',12),('Heading 1',15),('Heading 2',11.5)]:
        s=d.styles[n]; s.font.name='Arial'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0)
        s.font.bold=n!='Subtitle'; s.paragraph_format.space_before=Pt(12 if n!='Title' else 7); s.paragraph_format.space_after=Pt(7)
        s.font.underline=False
    for border in d.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    p=add(d,'Eswara Raj M'); p.runs[0].bold=True
    add(d,'CB.SC.U4CSE23022')
    d.add_paragraph(title,'Title'); add(d,'Gopher AI road damage detection project','Subtitle')
    add(d,'Evaluation date 15 October 2026')
    add(d,scope)
    f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    r=f.add_run('Page '); r.font.size=Pt(8)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); f._p.append(field)
    d.core_properties.author='Eswara Raj M'; d.core_properties.title=title
    d.core_properties.subject='Gopher AI evaluation for CB.SC.U4CSE23022'
    return d
def table(d,headers,rows,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,h in enumerate(headers): t.rows[0].cells[i].text=str(h)
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row): cells[i].text=str(x)
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+name); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
    t._tbl.tblPr.append(borders)
    repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for ri,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr(); margins=OxmlElement('w:tcMar')
            for edge,v in [('top','85'),('bottom','85'),('left','100'),('right','100')]:
                x=OxmlElement('w:'+edge); x.set(qn('w:w'),v); x.set(qn('w:type'),'dxa'); margins.append(x)
            pr.append(margins)
            sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'20374B' if ri==0 else ('F1F5F8' if ri%2 else 'FFFFFF')); pr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.05
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT if ci==0 else WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size=Pt(9); r.font.bold=ri==0
                    r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
    d.add_paragraph().paragraph_format.space_after=Pt(1)
def image(d,path,width=6.3):
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(path),width=Inches(width))
def refs(d,pages,extra):
    sub(d,'References and evidence')
    add(d,f'[1] CDAO Assessment and Assurance. Test and Evaluation of Artificial Intelligence Models. April 2024. Printed pages {pages}; the PDF page number is one greater than the printed number.')
    add(d,'[2] Gopher AI evaluation executed on 15 October 2026. '+extra)
def save(d,n):
    if len(sys.argv)>1 and sys.argv[1]!=n: return
    d.save(OUT/n); print(n)

# Explainability report: local probe, confidence reliability and subgroup limits.
d=new('Gopher AI XAI Report','This report explains the observed behaviour of gopher-ai_v1.pt through a local occlusion experiment, confidence reliability analysis and country subgroup results. The occlusion study identified a region whose removal suppressed a selected pothole detection. It supports local interpretation, but does not establish a complete explanation of the model.')
head(d,'Model and explanation objective')
add(d,'Gopher AI uses a supervised YOLO11n detector with 2,590,815 unfused parameters. It predicts bounding boxes, class labels and confidence scores for longitudinal crack, transverse crack, alligator crack, other corruption and pothole. The explanation objective is to make a road inspection prediction easier to review, especially when confidence alone could be misleading.')
head(d,'Local occlusion method')
add(d,'A reproducible random sample of 300 test images was selected with seed 23022. Among correctly matched pothole detections at confidence at least 0.25 and intersection over union at least 0.50, the highest-confidence detection was chosen. Its image was resized to a 512-pixel longest side and divided into a 6 by 6 grid.')
add(d,'Each of the 36 grid cells was replaced with gray pixel value 114 in a separate inference run. The tracked score was the highest remaining pothole confidence whose box overlapped the original predicted box at IoU at least 0.50; the score was zero if no qualifying box remained. Sensitivity equals original confidence minus the score after masking.')
it=X['interpretability']
table(d,['Measurement','Value'],[['Selected image',it['image']],['Original confidence',f"{it['baseline_confidence']:.4f}"],['Occlusion runs','36'],['Maximum confidence decrease',f"{it['max_confidence_drop']:.4f}"],['Mean signed decrease',f"{it['mean_confidence_drop']:.4f}"],['Minimum signed decrease',f"{it['min_confidence_drop']:.4f}"]],[4.4,2.3])
add(d,'The maximum decrease equals the original score: at least one mask eliminated the tracked detection under the matching rule. A negative decrease means the mask increased confidence. A score decrease can also reflect a localization change, not only weaker classification.')
page(d,'Visual explanation and its limits')
image(d,E/'report_figures/occlusion.png')
add(d,'Figure 1. Left: the selected correctly matched pothole prediction. Right: the confidence decrease when each cell is masked. Red indicates a larger decrease; blue indicates an increase. The colours show responses to separate masking experiments, not an attention map from the network.')
head(d,'Interpretation')
add(d,'The largest response lies around the selected detection region. Removing information there makes the tracked pothole detection disappear. This is evidence of local sensitivity to the displayed road region for this image and model configuration. It does not prove that all predictions rely on meaningful road-damage features.')
head(d,'Scope of the explanation')
add(d,'Occlusion is a post-hoc, model-agnostic technique. A gray patch can itself create an unfamiliar input, and a 6 by 6 grid provides coarse spatial resolution. The selected example is a strong correct prediction, so it cannot represent missed detections or typical explanation quality across the dataset.')
add(d,'No SHAP, LIME, Grad-CAM, causal attribution study or human explanation-usability assessment was performed. Those methods must not be inferred from this visualization. A broader explanation study should sample correct detections, false alarms and missed objects across classes and countries, using a selection rule fixed before reviewing the outputs.')
page(d,'Confidence reliability and subgroup evidence')
c=D['detection_calibration']
add(d,f"Across {c['predictions']:,} detections retained at confidence 0.25, expected calibration error was {c['ece_10_equal_width_bins']:.4f} and Brier score was {c['brier_score']:.4f}. Correctness required one-to-one same-class matching at IoU 0.50. Ten equal-width confidence bins were used, with empty bins contributing no weight.")
image(d,E/'report_figures/calibration.png',5.8)
add(d,'Figure 2. Average confidence compared with observed precision. The 0.90 to 1.00 bin contains only four detections, so the final downward point is highly uncertain. This analysis is conditional on retained detections and does not include the uncertainty of missed objects.')
table(d,['Country prefix','Recall','F1'],[[{'United':'United States','Czech':'Czech Republic'}.get(k,k),pc(v['recall']),pc(v['f1'])] for k,v in D['subgroups'].items()],[3.5,1.6,1.6])
add(d,'The country recall gap is 50.19 percentage points. Filename-derived country groups indicate possible domain differences, but class mix, scene conditions and annotation practices may also explain the gap. These results do not establish demographic fairness or a causal geographic effect.')
refs(d,'10 11 and 14','Evidence files: extended_results.json, analysis.json and the occlusion and calibration figures. The Testing Methodology report provides dataset and metric definitions.')
save(d,'CSE23022-XAI-REPORT.docx')

# Main methods and results, preserving primary metrics and data/operational caveats.
d=new('Gopher AI Testing Methodology','This report records the model evaluation design, data checks, full-test results and operational limits for gopher-ai_v1.pt. On 5,076 test images, mAP at IoU 0.50 was 49.04 percent and mAP averaged over IoU 0.50 to 0.95 was 24.50 percent. The evaluation follows the supplied CDAO framework at the model level.')
head(d,'Project objective and evaluation scope')
add(d,'The project detects five categories of road damage in dashcam imagery to support human inspection. A false alarm increases review work, while a missed defect can leave maintenance needs unrecorded. This evaluation concerns the frozen detector; it does not certify a complete camera, dashboard, vehicle or maintenance workflow.')
table(d,['Framework area','Method used','Evidence'],[
['Correctness','Full labeled test split','Detection AP precision recall and F1'],
['Data quality','Decode label and split audit','Counts defects and exact hashes'],
['Uncertainty','Image bootstrap and calibration','Intervals ECE and Brier scores'],
['Bias','Country subgroup comparison','Country precision recall and F1'],
['Robustness','Paired synthetic corruptions','300 images across nine conditions'],
['Interpretability','Local occlusion','36 masks on one pothole example'],
['Latency and resilience','Timed inference and restoration','100 timings and 20 sequences']],[1.3,2.6,2.8])
head(d,'Fixed experimental controls')
add(d,'Weights were unchanged. No retraining or threshold selection was performed. Full-set validation used input size 512, batch size 8, FP32, NMS IoU 0.70, maximum 300 detections and score floor 0.001 for AP. Fixed operating-point results use confidence 0.25 and same-class matching IoU 0.50. Testing used a GTX 1650 with four PyTorch CPU threads.')
add(d,'The framework is strategy-level guidance rather than a mandatory numerical pass standard. No stakeholder acceptance threshold was specified; results are reported without an overall pass or fail grade. The separate XAI and Ablation reports contain the detailed explanation and sensitivity studies.')
page(d,'Dataset and quality audit')
table(d,['Split','Images','Objects','Backgrounds','Background share'],[[k.title(),f"{v['images']:,}",f"{v['objects']:,}",f"{v['backgrounds']:,}",pc(v['backgrounds']/v['images'])] for k,v in A['splits'].items()],[1.1,1.2,1.3,1.4,1.7])
table(d,['Class','Train','Validation','Test'],[[k,f"{A['splits']['train']['classes'][str(i)]:,}",f"{A['splits']['val']['classes'][str(i)]:,}",f"{A['splits']['test']['classes'][str(i)]:,}"] for i,k in enumerate(D['per_class'])],[3.1,1.2,1.2,1.2])
add(d,'The object audit excludes a zero-width training annotation and retains duplicate rows while flagging them. Every test image decoded successfully and had a label file. No invalid test rows, duplicate test annotations, orphan test labels or exact duplicate test images were found.')
head(d,'Findings in development data')
add(d,'Training contains a zero-width box in Japan_001265 row 4, duplicate annotation rows in Japan_006916 and Japan_011427, and one excess exact duplicate image. Validation contains duplicate annotation rows in Japan_006536. The source data were not edited.')
head(d,'Split integrity and representativeness')
add(d,'All three pairwise split comparisons had zero shared filenames and zero shared image SHA-256 hashes. This rules out observed exact overlap, but not near-duplicate scenes or common routes. Earlier evaluation outputs exist, so historical test use for model selection cannot be excluded.')
add(d,'Training backgrounds account for 5.98 percent of images, versus 35.26 percent in testing. Country prefixes cover Czech Republic, India, Japan, Norway and United States. Camera, weather, night coverage and route metadata were not independently verified. Expert annotation correctness and deployment representativeness remain separate checks.')
add(d,'The project documents filtering drone and motorbike viewpoints, keeping 1,000 training background images, and resizing training and validation images. The audit verifies the current files rather than reconstructing every historical preprocessing action.')
page(d,'Detection and image classification results')
add(d,'For fixed-threshold detection, predictions are ordered by decreasing confidence and matched one-to-one to an unused ground-truth box of the same class at IoU at least 0.50. Unmatched predictions are false positives; unmatched annotations are false negatives.')
table(d,['Class','TP','FP','FN','Precision','Recall','F1'],[[k,v['tp'],v['fp'],v['fn'],pc(v['precision']),pc(v['recall']),pc(v['f1'])] for k,v in D['per_class'].items()]+[['All micro',D['overall']['tp'],D['overall']['fp'],D['overall']['fn'],pc(D['overall']['precision']),pc(D['overall']['recall']),pc(D['overall']['f1'])]],[2.05,.6,.6,.6,.95,.95,.95])
add(d,'Precision = TP / (TP + FP). Recall = TP / (TP + FN). F1 = 2PR / (P + R). Micro results aggregate object counts before calculating the ratios.')
table(d,['Class','AP at IoU 0.50','AP at IoU 0.50 to 0.95'],[[v['Class'],pc(v['mAP50']),pc(v['mAP50-95'])] for v in B['per_class']]+[['Mean across five classes',pc(B['metrics']['metrics/mAP50(B)']),pc(B['metrics']['metrics/mAP50-95(B)'])]],[3.1,1.6,2.0])
add(d,'AP integrates the precision-recall curve. The stricter AP measure also averages ten IoU thresholds. These AP values use the Ultralytics validation matching implementation and multi-label NMS; fixed-threshold counts use the matching rule stated above.')
head(d,'Secondary damage presence classification')
z=D['image_level']
add(d,f"An image is positive when it has any damage annotation; its score is the maximum detection confidence. At 0.25, TP={z['tp']:,}, FP={z['fp']:,}, TN={z['tn']:,} and FN={z['fn']:,}. Accuracy is {pc(z['accuracy'])}, ROC-AUC is {z['roc_auc']:.4f} and specificity is {pc(z['specificity'])}. This image-level task ignores location and damage category, so its accuracy is not detection accuracy.")
page(d,'Statistical uncertainty and runtime behaviour')
add(d,'Metric intervals were calculated from 2,000 image-level bootstrap samples with replacement and seed 23022. Each resampled image carries all its annotations and detections. This evaluates sampling variability for the fixed model; no model was retrained.')
table(d,['Metric','Estimate','95 percent percentile interval'],[[label,pc(val),' to '.join(pc(x) for x in D['bootstrap_95ci'][key])] for key,label,val in [('precision','Detection precision',D['overall']['precision']),('recall','Detection recall',D['overall']['recall']),('f1','Detection F1',D['overall']['f1']),('image_accuracy','Image accuracy',z['accuracy'])]],[2.6,1.3,2.8])
add(d,'The intervals assume exchangeable images. Correlated frames or routes may make them too narrow. They do not include annotation error, training variability or uncertainty on unseen environments. Image-level Brier score is 0.1907; maximum detection confidence is only a proxy for image probability.')
head(d,'Latency measurements')
l=X['latency']; r=X['resilience']
table(d,['Measure','Value'],[['Warmup and timed images','10 and 100'],['Mean and median prediction call',f"{l['mean_ms']:.2f} ms and {l['median_ms']:.2f} ms"],['95th and 99th percentile',f"{l['p95_ms']:.2f} ms and {l['p99_ms']:.2f} ms"],['Minimum and maximum',f"{l['min_ms']:.2f} ms and {l['max_ms']:.2f} ms"],['Serial throughput',f"{l['serial_images_per_second']:.2f} images per second"],['Peak CUDA tensor allocation',f"{l['peak_cuda_allocated_bytes']/1024**2:.2f} MiB"]],[3.9,2.8])
add(d,'Batch-one timing synchronizes CUDA around the complete prediction call at 512 by 512 resolution and confidence 0.25. It includes API overhead, preprocessing, inference and postprocessing. File decoding, initial resizing, camera capture, network transport and display are excluded. Tensor allocation excludes allocator reserve and other processes. The batch-eight validator separately measured 6.92 ms of inference per image, an amortized component time.')
head(d,'Recovery probe')
add(d,f"All {r['restored_output_trials']} of {r['trials']} clean-darkened-clean sequences restored their initial predictions. The mean restoration-check time was {r['mean_recovery_call_ms']:.2f} ms, including output conversion and comparison. Classes were equal; confidence and box absolute tolerances were 1e-6 and 1e-4 with relative tolerance 1e-5. This tests lingering input effects in a stateless detector, not service crashes or sensor recovery.")
page(d,'Reproducibility and evaluation limits')
table(d,['Model fact','Recorded value'],[['Checkpoint','gopher-ai_v1.pt'],['Architecture and classes','YOLO11n and five road damage classes'],['Unfused parameters','2,590,815'],['Checkpoint size','21,254,689 bytes'],['Checkpoint epoch','78 zero based corresponding to epoch 79'],['Training settings','AdamW initial LR 0.0015 batch 16 image size 512'],['Training controls','Maximum 100 epochs patience 10 seed 0 AMP disabled'],['Evaluation software','Python 3.12.10 PyTorch 2.11.0+cu128 Ultralytics 8.4.90'],['Evaluation device','NVIDIA GeForce GTX 1650 with 4 GiB']],[2.3,4.4])
add(d,'Checkpoint SHA-256: '+A['model_sha256'])
add(d,'Image and label manifest SHA-256: '+A['manifest_sha256'])
add(d,'Reproduce the measurements from the project root with evaluation/audit_data.py, run_baseline.py, analyze_results.py and run_extended.py. The output/evaluation folder preserves the audit, metadata, metrics, full predictions, per-image CSV and manifest. The checkpoint matches the local best.pt by SHA-256.')
head(d,'Remaining evaluation work')
add(d,'No training-fold cross-validation, independent comparator experiment, adversarial red teaming or data-poisoning test was performed. Temporal drift requires newer timestamped production data. Demographic fairness, legal compliance, user explanation quality and complete application resilience were not established.')
add(d,'At the chosen threshold, 52.10 percent of annotated objects are missed and 18.66 percent of background images receive at least one alarm. Priorities are reviewing missed potholes and weak subgroups, correcting verified label defects, selecting thresholds on separate validation data, and collecting route-disjoint deployment images. No autonomous deployment or real-time pass claim follows from these tests.')
refs(d,'3 9 13 to 17 19 to 28 31 to 37 and 59 to 63','Sources include the original 11-page evaluation report, audit.json, baseline_metrics.json, analysis.json, extended_results.json, model_metadata.json and project training scripts.')
save(d,'CSE23022-Testing-methodology.docx')

# Architecture and training report documents the implemented detector.
d=new('Gopher AI Model Architecture and Training Report','Gopher AI uses YOLO11n to locate and classify road damage in dashcam images. This report documents the implemented detector, dataset preparation, recorded training configuration and checkpoint evidence. It complements the XAI, Testing Methodology and Ablation reports by explaining how the evaluated model was built.')
head(d,'Model purpose and outputs')
add(d,'The supervised detection task combines localization with damage classification. Each retained prediction contains a rectangular box, a class label and a confidence score. The five labels are longitudinal crack, transverse crack, alligator crack, other corruption and pothole. These image-space predictions support inspection; they do not directly measure defect depth, physical area or repair priority.')
table(d,['Model property','Verified value'],[['Evaluated checkpoint','gopher-ai_v1.pt'],['Architecture scale','YOLO11n nano'],['Unfused parameter count','2,590,815'],['Checkpoint file size','21,254,689 bytes'],['Recorded training input size','512 pixels'],['Number of damage classes','5']],[3.3,3.4])
head(d,'Architecture and information flow')
add(d,'An input image passes through learned feature extraction, fusion of features at different resolutions and a detection head. The inspected model configuration contains convolutional layers, C3k2 blocks, SPPF and C2PSA in the backbone. Upsampling and concatenation combine feature maps in the head before the Detect module consumes three feature levels.')
table(d,['Stage','Observed components','Role in this detector'],[['Feature extraction','Conv and C3k2','Learn image features at successive resolutions'],['Backbone processing','SPPF and C2PSA','Process the deeper image representation'],['Feature fusion','Upsample Concat and C3k2','Combine information across feature levels'],['Prediction','Detect with five classes','Produce localization and class predictions']],[1.35,2.35,3.0])
add(d,'The component list describes the loaded architecture. It does not measure how much any individual block improves performance; that requires the controlled structural experiments described in the Ablation Report.')
page(d,'Dataset preparation and training configuration')
add(d,'The local dataset is a derivative of RDD2022 with YOLO-format box annotations. A nonempty label row records class ID, normalized box centre and normalized width and height. Images without annotated damage are retained as negative examples. The current split audit found the following quantities.')
table(d,['Split','Images','Valid annotation rows','Background images'],[[k.title(),f"{v['images']:,}",f"{v['objects']:,}",f"{v['backgrounds']:,}"] for k,v in A['splits'].items()],[1.1,1.4,2.2,2.0])
add(d,'The audit excludes one zero-width training row from its valid count and flags duplicate annotation rows separately. Exact image hashes and filenames did not overlap between splits, although near-duplicate scenes and historical test-set use remain unverified.')
head(d,'Recorded preprocessing pipeline')
add(d,'Afilter_dash.py targets China_Drone and China_MotorBike files across the splits. prune_backgrounds.py keeps the first 1,000 empty-label training pairs encountered and moves excess pairs into a backup folder. resize_images.py reduces training and validation JPEGs larger than 640 pixels on their longest side, preserving aspect ratio and saving at JPEG quality 85. Test images are not resized by that script.')
add(d,'These are the behaviours of the project scripts; they were not rerun to create this report. Reducing background samples changes the training distribution: backgrounds are 5.98 percent of training images and 35.26 percent of test images. This difference matters when interpreting false alarms.')
head(d,'Training settings from the checkpoint')
table(d,['Setting','Recorded value'],[['Optimizer and initial learning rate','AdamW and 0.0015'],['Weight decay and final LR factor','0.0005 and 0.01'],['Maximum epochs and early stopping patience','100 and 10'],['Batch size and image size','16 and 512'],['Loader workers and cache','3 and enabled'],['Automatic mixed precision','Disabled'],['Random seed and deterministic flag','0 and enabled'],['Warmup duration','3 epochs']],[3.7,3.0])
add(d,'new_train.py initializes yolo11n.pt and starts detector training. resume_train.py resumes the saved train-10 last.pt checkpoint. The evaluated checkpoint records a resumed run. The settings document that run; training was not repeated during the evaluation.')
page(d,'Augmentation checkpoint selection and results')
head(d,'Recorded augmentation settings')
table(d,['Augmentation setting','Value'],[['Horizontal flip probability','0.5'],['Vertical flip and rotation','0.0 and 0.0 degrees'],['Mosaic probability and closing period','1.0 and final 10 scheduled epochs'],['Scale and translation','0.5 and 0.1'],['HSV hue saturation and value','0.015 0.7 and 0.4'],['Mixup and cutmix','Both 0.0']],[3.7,3.0])
add(d,'The configuration records no vertical flips or rotation and enables horizontal flipping. Other settings vary colour, scale and image composition. Their individual effects were not isolated. A shared options file also contains settings for other tasks; their presence alone does not show they were applied to detection training.')
head(d,'Saved model and training history')
add(d,'The supplied checkpoint has zero-based epoch 78, corresponding to epoch 79, and a saved timestamp of 18 September 2026 at 10:50:16 without a recorded timezone. Its SHA-256 matches the local train-10 best.pt. The training CSV includes rows through epoch 80. A 100-epoch limit in the configuration does not establish that 100 epochs completed or why the recorded run ended.')
table(d,['Evidence source','mAP at IoU 0.50','mAP at IoU 0.50 to 0.95'],[['Training log validation row at epoch 79','50.43%','25.66%'],['Training log validation row at epoch 80','50.46%','25.66%'],['Fresh checkpoint test on 5,076 images',pc(B['metrics']['metrics/mAP50(B)']),pc(B['metrics']['metrics/mAP50-95(B)'])]],[3.1,1.6,2.0])
add(d,'The validation rows describe historical development measurements; the test row comes from the fresh evaluation. These datasets and evaluation contexts differ, so the gap alone does not establish overfitting. Detailed per-class performance, uncertainty and latency are in the Testing Methodology report.')
add(d,'Checkpoint SHA-256: '+A['model_sha256'])
sub(d,'References and reproducibility')
add(d,'[1] Local project sources: new_train.py, resume_train.py, Afilter_dash.py, prune_backgrounds.py, resize_images.py, rdd2022/data.yaml and runs/detect/train-10/args.yaml and results.csv.')
add(d,'[2] Inspected checkpoint architecture and evaluation records: gopher-ai_v1.pt, model_metadata.json, audit.json and baseline_metrics.json. Evaluation date 15 October 2026.')
add(d,'[3] CDAO Assessment and Assurance. Test and Evaluation of Artificial Intelligence Models. April 2024. Printed pages 31 to 37, 41 to 49 and 62 to 63, covering data preparation, model lifecycle and documentation.')
save(d,'CSE23022-Model-Architecture-and-Training-Report.docx')

# Controlled sensitivity experiments and a clearly separate structural ablation plan.
d=new('Gopher AI Ablation and Sensitivity Report','This report separates completed input and threshold sensitivity experiments from structural model ablations. The completed experiments kept gopher-ai_v1.pt fixed. No network module was removed and retrained, so the results do not quantify the contribution of individual YOLO11n components.')
head(d,'Experimental design')
add(d,'A uniform random sample of 300 test images was chosen without replacement using seed 23022. The same 481 annotated objects and images were reused across conditions. Country counts were Czech Republic 19, India 60, Japan 94, Norway 83 and United States 44. Each source image was resized to a 512-pixel longest side before transformation, followed by square letterboxing.')
add(d,'Inference used FP32, image size 512, single-label prediction NMS at IoU 0.70, maximum 300 detections, and a score floor of 0.001. Reported detection results retained scores at least 0.25 and used one-to-one same-class matches at IoU 0.50. Labels were unchanged except that horizontal-flip box centres were mirrored.')
labels={'clean':'Clean reference','dark_0.75':'Brightness times 0.75','dark_0.50':'Brightness times 0.50','bright_1.50':'Brightness times 1.50','gaussian_blur_5':'Gaussian blur 5 by 5','gaussian_noise_15':'Gaussian noise sigma 15','jpeg_quality_30':'JPEG quality 30','dark_0.50_and_blur_5':'Darkness and blur','horizontal_flip':'Horizontal flip'}
table(d,['Condition','Precision','Recall','F1','F1 change pp'],[[labels[k],pc(v['precision']),pc(v['recall']),pc(v['f1']),f"{v.get('f1_delta_percentage_points',0):+.2f}"] for k,v in X['corruptions'].items()],[2.5,1.05,1.05,1.05,1.05])
add(d,'Brightness values were multiplied and clipped to 0 to 255. Blur used a 5 by 5 kernel with sigma 1.5. Gaussian noise had standard deviation 15 in the 0 to 255 pixel range. JPEG was encoded at quality 30 and decoded. Percentage-point changes are relative to the clean subset, whose F1 was 45.39 percent.')
page(d,'Factor interactions and prediction consistency')
head(d,'Darkness and blur comparison')
add(d,'Clean, darkness only, blur only, and darkness with blur form a complete two by two factorial design for these two binary factors. This covers the pair of tested factors, not every pair of environmental conditions.')
table(d,['Brightness factor','Blur','F1'],[['1.00','Off',pc(X['corruptions']['clean']['f1'])],['0.50','Off',pc(X['corruptions']['dark_0.50']['f1'])],['1.00','On',pc(X['corruptions']['gaussian_blur_5']['f1'])],['0.50','On',pc(X['corruptions']['dark_0.50_and_blur_5']['f1'])]],[2.3,2.1,2.3])
add(d,'Blur reduced F1 by 19.98 percentage points. Darkness with blur produced the lowest F1, 23.90 percent, a decrease of 21.50 points. Darkness alone increased F1 by 1.02 points on this sample; this is not evidence that darkness generally improves detection. No paired confidence interval or significance test was calculated for these differences.')
head(d,'Horizontal reflection relation')
mt=X['metamorphic']
add(d,'The metamorphic hypothesis was that mirroring an image should approximately preserve damage detections after coordinates are reflected back. Same-class boxes were matched at IoU at least 0.50 after applying the reverse coordinate transform.')
table(d,['Consistency measure','Value'],[['Matched boxes','312'],['Clean and flipped retained boxes','387 and 375'],['Clean detection retention',pc(mt['clean_detection_retention'])],['Symmetric consistency F1',pc(mt['symmetric_consistency_f1'])],['Images with complete consistency','199 of 300']],[4.3,2.4])
add(d,'Complete consistency includes images with no detections in either run. The relation measures output stability, not correctness by itself. Reflected scenes may also change directional semantics. Ground-truth detection F1 and prediction consistency should therefore be interpreted together.')
head(d,'Limits of these comparisons')
add(d,'The full-set baseline and subset results use different sample sizes and inference preprocessing. Compare each corruption with its own clean subset result. Synthetic blur, darkness and noise do not independently establish performance in real rain, night driving or new camera conditions.')
page(d,'Threshold study and future structural ablations')
head(d,'Measured confidence threshold sensitivity')
add(d,'The full test predictions were rescored at four thresholds using the same matching rule. Weights and data were fixed. The table describes the precision and recall trade-off; it was not used to select an optimal threshold on the test set.')
table(d,['Threshold','Precision','Recall','F1','FP','FN'],[[f"{v['confidence']:.2f}",pc(v['precision']),pc(v['recall']),pc(v['f1']),v['fp'],v['fn']] for v in D['threshold_sweep']],[1.05,1.15,1.15,1.15,1.1,1.1])
head(d,'Structural ablations not yet executed')
add(d,'To measure a component contribution, train an independent baseline and a variant with that component removed or changed under matched data splits, training budgets and seeds. Repeating across seeds is necessary to distinguish a component effect from training variability. The following are proposed experiments, not measured results.')
table(d,['Candidate experiment','Comparison','Current result'],[['Training augmentation','Recorded recipe versus one augmentation disabled','Not measured'],['Input resolution','512 versus another chosen training size','Not measured'],['Attention module','Standard YOLO11n versus a validated architecture without the selected module','Not measured'],['Background sampling','Current pruning versus a defined alternative training distribution','Not measured']],[1.7,3.6,1.4])
add(d,'Each variant should be tuned on validation data and evaluated on the same untouched real-image test set. Report class AP, precision, recall, F1, parameter count and latency with the same hardware and measurement boundaries. Do not substitute inference-time corruption results for module-removal results.')
refs(d,'12 20 22 23 and 45 to 46','Evidence files: extended_results.json and analysis.json. The Testing Methodology report defines the full dataset and operating point. The XAI report contains the separate occlusion probe.')
save(d,'CSE23022_Ablation_Report.docx')
